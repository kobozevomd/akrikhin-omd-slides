# Выгрузка структуры исходной деки для пересборки: по слайду — тексты с координатами (px при ширине 1920),
# графики с данными, таблицы, картинки (в media/), заметки.
#   python3 tools/extract_deck.py <deck.pptx> <out_dir>
# → out_dir/slide-NN.txt (читать глазами), slide-NN.json (данные), media/img_<hash>.<ext>
import sys, os, json, hashlib
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def main(src, out):
    os.makedirs(os.path.join(out, 'media'), exist_ok=True)
    p = Presentation(src)
    K = 1920 / p.slide_width

    def walk(shapes, acc, depth=0, gx=None):
        for sh in shapes:
            try:
                x, y, w, h = [round(v * K) if v is not None else None for v in (sh.left, sh.top, sh.width, sh.height)]
            except Exception:
                x = y = w = h = None
            if gx and x is not None:
                # координаты детей группы — в системе группы; переводим в координаты слайда
                (ox, oy, sx, sy, cx0, cy0) = gx
                x = round(ox + (x - cx0) * sx); y = round(oy + (y - cy0) * sy)
                w = round(w * sx); h = round(h * sy)
            item = {'t': str(sh.shape_type).split('.')[-1].split(' ')[0] if sh.shape_type else 'SHAPE', 'name': sh.name,
                    'x': x, 'y': y, 'w': w, 'h': h, 'd': depth}
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                acc.append(item)
                try:
                    ch = sh._element.xpath('./p:grpSpPr/a:xfrm')[0]
                    off = ch.find('{http://schemas.openxmlformats.org/drawingml/2006/main}off')
                    ext = ch.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ext')
                    choff = ch.find('{http://schemas.openxmlformats.org/drawingml/2006/main}chOff')
                    chext = ch.find('{http://schemas.openxmlformats.org/drawingml/2006/main}chExt')
                    sx = int(ext.get('cx')) / max(1, int(chext.get('cx')))
                    sy = int(ext.get('cy')) / max(1, int(chext.get('cy')))
                    g = (x, y, sx, sy, round(int(choff.get('x')) * K), round(int(choff.get('y')) * K))
                except Exception:
                    g = None
                walk(sh.shapes, acc, depth + 1, g)
                continue
            if getattr(sh, 'has_text_frame', False) and sh.has_text_frame and sh.text_frame.text.strip():
                paras = []
                for para in sh.text_frame.paragraphs:
                    t = ''.join(r.text for r in para.runs).strip()
                    if not t:
                        continue
                    b = any(r.font.bold for r in para.runs)
                    sz = next((r.font.size.pt for r in para.runs if r.font.size), None)
                    col = None
                    for r in para.runs:
                        try:
                            if r.font.color and r.font.color.type is not None and r.font.color.rgb is not None:
                                col = str(r.font.color.rgb); break
                        except Exception:
                            pass
                    paras.append(('**' if b else '') + t + ('**' if b else '') + (f' ·{sz:g}pt' if sz else '') + (f' #{col}' if col else ''))
                item['text'] = paras
            try:
                if sh.fill.type == 1:
                    item['fill'] = str(sh.fill.fore_color.rgb)
            except Exception:
                pass
            if getattr(sh, 'has_chart', False) and sh.has_chart:
                ch = sh.chart
                cd = {'type': str(ch.chart_type).split('.')[-1].split(' ')[0], 'plots': []}
                for pl in ch.plots:
                    cats = [str(c) for c in pl.categories]
                    cd['plots'].append({'type': str(pl.__class__.__name__), 'cats': cats,
                                        'series': [{'name': s.name, 'values': list(s.values)} for s in pl.series]})
                item['chart'] = cd
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    blob = sh.image.blob
                    hsh = hashlib.md5(blob).hexdigest()[:10]
                    fn = f'img_{hsh}.{sh.image.ext}'
                    fp = os.path.join(out, 'media', fn)
                    if not os.path.exists(fp):
                        open(fp, 'wb').write(blob)
                    item['image'] = fn
                except Exception as e:
                    item['image'] = f'ERR {e}'
            if getattr(sh, 'has_table', False) and sh.has_table:
                item['table'] = [[c.text for c in r.cells] for r in sh.table.rows]
            acc.append(item)

    for i, s in enumerate(p.slides, 1):
        acc = []
        walk(s.shapes, acc)
        notes = ''
        try:
            notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ''
        except Exception:
            pass
        data = {'slide': i, 'layout': s.slide_layout.name, 'shapes': acc, 'notes': notes}
        json.dump(data, open(os.path.join(out, f'slide-{i:02d}.json'), 'w'), ensure_ascii=False, indent=1)
        lines = [f'# Слайд {i} (layout: {s.slide_layout.name})']
        for it in acc:
            ind = '  ' * it['d']
            head = f"{ind}[{it['t']}] {it['x']},{it['y']} {it['w']}x{it['h']}" + (f" fill={it['fill']}" if 'fill' in it else '')
            if 'text' in it:
                lines.append(head + ' :: ' + ' ¶ '.join(it['text']))
            elif 'chart' in it:
                lines.append(head + ' CHART ' + json.dumps(it['chart'], ensure_ascii=False))
            elif 'image' in it:
                lines.append(head + ' IMAGE ' + it['image'])
            elif 'table' in it:
                lines.append(head + ' TABLE ' + json.dumps(it['table'], ensure_ascii=False))
            elif it['t'] == 'GROUP':
                lines.append(head + ' GROUP')
        if notes:
            lines.append('NOTES: ' + notes)
        open(os.path.join(out, f'slide-{i:02d}.txt'), 'w').write('\n'.join(lines))
        print(i, s.slide_layout.name, 'chars', sum(len(' '.join(a.get('text', []))) for a in acc),
              'charts', sum('chart' in a for a in acc), 'imgs', sum('image' in a for a in acc), 'tables', sum('table' in a for a in acc))


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
