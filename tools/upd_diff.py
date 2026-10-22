#!/usr/bin/env python3
"""Что поменял клиент: сравнение моей RU-деки с файлом, который вернул Роман.

    python3 tools/upd_diff.py <моя RU.pptx> <правленая RU.pptx>

По каждому слайду: изменения текста (строки −/+), сдвинутые или изменённые по размеру фигуры, скрытие.
Неразрывные пробелы и пробелы схлопываются — типографика не считается правкой.
Высота таблицы, растянутая самим PowerPoint при открытии (строки те же), отмечается как «таблица (PowerPoint)» —
её не переносим.
"""
import difflib, re, sys
from pptx import Presentation


def items(sl):
    out = []

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == 6:
                walk(sh.shapes)
                continue
            if sh.has_text_frame:
                t = re.sub(r'\s+', ' ', sh.text_frame.text.replace(' ', ' ')).strip()
                if t:
                    out.append(t)
            if getattr(sh, 'has_table', False) and sh.has_table:
                for r in sh.table.rows:
                    for c in r.cells:
                        t = re.sub(r'\s+', ' ', c.text.replace(' ', ' ')).strip()
                        if t:
                            out.append(t)
    walk(sl.shapes)
    return out


def main():
    A, B = Presentation(sys.argv[1]), Presentation(sys.argv[2])
    K = 1920 / A.slide_width
    px = lambda v: None if v is None else round(v * K)
    print(f'слайдов: {len(A.slides)} → {len(B.slides)}')
    any_change = False
    for i, (a, b) in enumerate(zip(A.slides, B.slides), 1):
        d = [l for l in difflib.unified_diff(items(a), items(b), lineterm='', n=0) if not l.startswith(('---', '+++', '@@'))]
        geo = []
        for x, y in zip(a.shapes, b.shapes):
            gx = (px(x.left), px(x.top), px(x.width), px(x.height))
            gy = (px(y.left), px(y.top), px(y.width), px(y.height))
            if gx != gy:
                tag = ''
                if getattr(x, 'has_table', False) and x.has_table and gx[:3] == gy[:3]:
                    rows_a = [px(r.height) for r in x.table.rows]
                    rows_b = [px(r.height) for r in y.table.rows]
                    if rows_a == rows_b:
                        tag = '  — таблица (PowerPoint растянул при открытии, не переносим)'
                txt = (y.text_frame.text[:40] if y.has_text_frame else '').replace('\n', ' ')
                geo.append(f'{y.name} «{txt}»: {gx} → {gy}{tag}')
        hid = (a._element.get('show'), b._element.get('show'))
        if d or geo or hid[0] != hid[1] or len(a.shapes) != len(b.shapes):
            any_change = True
            print(f'\n=== слайд {i}')
            for l in d:
                print('  ', l[:400])
            for g in geo:
                print('   geo', g)
            if len(a.shapes) != len(b.shapes):
                print(f'   фигур: {len(a.shapes)} → {len(b.shapes)}')
            if hid[0] != hid[1]:
                print(f'   скрытие: {hid[0]} → {hid[1]}')
    if not any_change:
        print('отличий нет')


if __name__ == '__main__':
    main()
