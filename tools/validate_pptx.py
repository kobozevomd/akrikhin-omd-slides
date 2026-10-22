#!/usr/bin/env python3
"""Проверка собранного .pptx перед отдачей: откроется ли файл в PowerPoint без «восстановления».

    python3 tools/validate_pptx.py <deck.pptx>

Проверяет:
  - архив читается, имена частей не повторяются;
  - каждая XML-часть и .rels разбираются (lxml);
  - у каждой части есть тип в [Content_Types].xml (Override или Default по расширению);
  - внутренние связи (.rels) ведут на существующие части;
  - все слайды из presentation.xml на месте;
  - на слайдах нет повторяющихся id фигур (cNvPr) и <a:pPr> стоит первым в абзаце;
  - python-pptx открывает файл и проходит по фигурам всех слайдов.
Код выхода 0 и строка «All validations PASSED!» — можно отдавать; иначе список ошибок.
"""
import posixpath, re, sys, zipfile
from lxml import etree

NS = {
    'ct': 'http://schemas.openxmlformats.org/package/2006/content-types',
    'r': 'http://schemas.openxmlformats.org/package/2006/relationships',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
}
RID = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'


def rels_base(rels_name):
    # ppt/slides/_rels/slide1.xml.rels → ppt/slides
    d = posixpath.dirname(posixpath.dirname(rels_name))
    return d


def main(path):
    errors, warnings = [], []
    try:
        z = zipfile.ZipFile(path)
    except zipfile.BadZipFile as e:
        print(f'FAILED: не zip-архив ({e})')
        return 1
    names = z.namelist()
    if len(names) != len(set(names)):
        errors.append('в архиве повторяются имена частей')
    parts = set(n for n in names if not n.endswith('/'))

    trees = {}
    for n in parts:
        if n.endswith('.xml') or n.endswith('.rels'):
            try:
                trees[n] = etree.fromstring(z.read(n))
            except etree.XMLSyntaxError as e:
                errors.append(f'{n}: XML не разбирается ({e})')

    ct = trees.get('[Content_Types].xml')
    if ct is None:
        errors.append('нет [Content_Types].xml')
    else:
        defaults = {e.get('Extension').lower() for e in ct.findall('ct:Default', NS)}
        overrides = {e.get('PartName').lstrip('/') for e in ct.findall('ct:Override', NS)}
        for n in parts:
            if n == '[Content_Types].xml':
                continue
            ext = n.rsplit('.', 1)[-1].lower() if '.' in n else ''
            if n not in overrides and ext not in defaults:
                errors.append(f'{n}: нет типа в [Content_Types].xml')
        orphans = sorted(o for o in overrides if o not in parts)
        if orphans:
            # pptxgenjs пишет лишние Override на slideMasterN; PowerPoint их пропускает, lib/postprocess.js вычищает
            warnings.append(f'[Content_Types].xml: {len(orphans)} Override на несуществующие части (не мешает), напр. {orphans[0]}')

    for n, t in trees.items():
        if not n.endswith('.rels'):
            continue
        base = rels_base(n)
        for rel in t.findall('r:Relationship', NS):
            if rel.get('TargetMode') == 'External':
                continue
            tgt = rel.get('Target')
            full = posixpath.normpath(posixpath.join(base, tgt)) if not tgt.startswith('/') else tgt.lstrip('/')
            if full not in parts:
                errors.append(f'{n}: связь {rel.get("Id")} ведёт на отсутствующую часть {full}')

    pres = trees.get('ppt/presentation.xml')
    prels = trees.get('ppt/_rels/presentation.xml.rels')
    n_slides = 0
    if pres is not None and prels is not None:
        rmap = {r.get('Id'): r.get('Target') for r in prels.findall('r:Relationship', NS)}
        for s in pres.findall('p:sldIdLst/p:sldId', NS):
            n_slides += 1
            tgt = rmap.get(s.get(RID))
            if not tgt or posixpath.normpath(posixpath.join('ppt', tgt)) not in parts:
                errors.append(f'presentation.xml: слайд {s.get("id")} без части ({tgt})')

    for n, t in trees.items():
        if not re.match(r'ppt/(slides|slideLayouts|slideMasters)/[^/]+\.xml$', n):
            continue
        ids = [e.get('id') for e in t.iter('{%s}cNvPr' % NS['p'])]
        dup = sorted({i for i in ids if ids.count(i) > 1})
        if dup:
            errors.append(f'{n}: повторяются id фигур {dup[:5]}')
        for p in t.iter('{%s}p' % NS['a']):
            kids = [c for c in p if isinstance(c.tag, str)]
            for k, c in enumerate(kids):
                if c.tag == '{%s}pPr' % NS['a'] and k != 0:
                    errors.append(f'{n}: <a:pPr> не первым в абзаце')
                    break
            else:
                continue
            break

    try:
        from pptx import Presentation
        prs = Presentation(path)
        for sl in prs.slides:
            for sh in sl.shapes:
                _ = sh.shape_type
    except Exception as e:  # noqa: BLE001 — любая ошибка открытия = файл битый
        errors.append(f'python-pptx не открывает файл: {e}')

    for w in warnings:
        print('  предупреждение:', w)
    if errors:
        print(f'FAILED: {len(errors)} ошибок')
        for e in errors[:50]:
            print('  -', e)
        return 1
    print(f'{path}: слайдов {n_slides}, частей {len(parts)}')
    print('All validations PASSED!')
    return 0


if __name__ == '__main__':
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(sys.argv[1]))
